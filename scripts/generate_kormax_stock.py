import json
import os
import re
from datetime import timedelta
from io import BytesIO

from lxml import etree
from openpyxl import load_workbook

from generate_home_food_stock import (
    PRODUCT_UPDATE_URL,
    api_json,
    download,
    fetch_orders,
    format_api_time,
    now_kyiv,
    parse_api_time,
)


CATALOG_XML_URL = (
    "https://catpaws.com.ua/content/export/"
    "3e9c244f28ee6d1e572f92646e76f6bb.xml"
)
KORMAX_XLSX_URL = "https://b2b.kormaxtrade.com.ua/feeds/edq0mnq846.xlsx"

OUTPUT_FILE = "public/kormax-stock.yml"
STATE_FILE = "public/kormax-stock-state.json"

# У замовленні №2506 товари вручну переведені на склад
# "Kormaxtrade - Білогородка". Воно використовується тільки для
# визначення числового ID складу. Під час першого запуску його продажі
# записуються у початковий знімок і повторно не віднімаються.
BOOTSTRAP_ORDER_ID = 2506

ALLOWED_BRANDS = {
    "Nature's Protection",
    "Nature's Protection Superior Care",
    "Nature's Protection Lifestyle",
    "Nature's Protection Prime",
    "Tauro Pro Line",
}

DEFAULT_TARGET = 20
DEFAULT_THRESHOLD = 5
MEDIUM_TARGET = 10
MEDIUM_THRESHOLD = 2
LOW_TARGET = 5
LOW_THRESHOLD = 1
CAT_CAN_TARGET = 60
CAT_CAN_THRESHOLD = 24
DOG_CAN_TARGET = 40
DOG_CAN_THRESHOLD = 16

BOX_PACK_SIZES_BY_COMPONENT = {}


def clean(value):
    return " ".join(str(value or "").split())


def normalize(value):
    return clean(value).casefold().replace("’", "'")


def load_site_catalog():
    root = etree.fromstring(download(CATALOG_XML_URL))
    products = {}

    for offer in root.xpath(".//offer"):
        sku = clean(offer.findtext("vendorCode"))
        name = clean(offer.findtext("name"))
        if sku and name:
            products[sku] = name

    return products


def load_supplier():
    workbook = load_workbook(
        BytesIO(download(KORMAX_XLSX_URL)),
        read_only=True,
        data_only=True,
    )
    sheet = workbook.active
    rows = sheet.iter_rows(values_only=True)
    headers = [clean(value) for value in next(rows)]
    columns = {name: index for index, name in enumerate(headers)}
    required = {"Артикул", "Назва", "Бренд", "Статус", "Кількість"}

    missing = required - set(columns)
    if missing:
        raise RuntimeError(
            "Kormax XLSX does not contain columns: " + ", ".join(sorted(missing))
        )

    products = {}
    for row in rows:
        sku = clean(row[columns["Артикул"]])
        brand = clean(row[columns["Бренд"]])
        if not sku or brand not in ALLOWED_BRANDS:
            continue

        products[sku] = {
            "name": clean(row[columns["Назва"]]),
            "brand": brand,
            "status": clean(row[columns["Статус"]]),
            "quantity_band": clean(row[columns["Кількість"]]),
        }

    return products


def parse_pack_size(name):
    match = re.search(
        r"(?:\*|×|x|х)\s*(\d{1,3})\s*(?:шт\.?|од\.?|pcs)?\b",
        normalize(name),
    )
    return int(match.group(1)) if match else None


def is_cat_product(name):
    text = normalize(name)
    return any(word in text for word in ("кот", "кошен", "cats", "cat ", "kitten"))


def is_dog_product(name):
    text = normalize(name)
    return any(word in text for word in ("собак", "цуцен", "dogs", "dog ", "puppy"))


def is_wet_food(name):
    text = normalize(name)
    return any(
        word in text
        for word in ("волог", "консерв", "пауч", "wet food")
    )


def build_product_catalog(site_catalog, supplier):
    products = {}
    boxes = {}

    # Беремо тільки товари, які одночасно є у прайсі постачальника
    # та вже заведені в каталозі CatPaws/SalesDrive.
    for sku, supplier_item in supplier.items():
        if sku in site_catalog:
            products[sku] = site_catalog[sku]

    # Власні коробки постачальник не передає. Знаходимо їх у каталозі
    # за суфіксом _box та рахуємо кількість із залишку одиничного SKU.
    for box_sku, box_name in site_catalog.items():
        if not box_sku.lower().endswith("_box"):
            continue

        component_sku = box_sku[:-4]
        if component_sku not in products:
            continue

        pack_size = parse_pack_size(box_name)
        if not pack_size:
            print(f"Skipped box without pack size: {box_sku} — {box_name}")
            continue

        products[box_sku] = box_name
        boxes[box_sku] = (component_sku, pack_size)

    BOX_PACK_SIZES_BY_COMPONENT.clear()
    BOX_PACK_SIZES_BY_COMPONENT.update(
        {component_sku: pack_size for component_sku, pack_size in boxes.values()}
    )

    return products, boxes


def supplier_available(item):
    status = normalize(item.get("status"))
    return "немає" not in status and "нет в наличии" not in status


def supplier_band(item):
    if not supplier_available(item):
        return "zero"

    band = normalize(item.get("quantity_band"))
    if "до 5" in band:
        return "low"
    if "6 до 10" in band:
        return "medium"
    if "11 до 100" in band or "від 101" in band or "от 101" in band:
        return "high"

    # Якщо постачальник позначив товар наявним, але не передав діапазон,
    # безпечніше показати мінімальний залишок, а не 20/60.
    return "low"


def stock_policy(sku, name, supplier):
    band = supplier_band(supplier[sku])
    if band == "zero":
        return 0, 0, band
    if band == "low":
        return LOW_TARGET, LOW_THRESHOLD, band
    if band == "medium":
        return MEDIUM_TARGET, MEDIUM_THRESHOLD, band

    # Великі залишки вологого корму дають 5 повних коробок:
    # котячі паучі/консерви — 5 × 12 = 60;
    # собачі — 5 × 8 = 40.
    if is_wet_food(name):
        pack_size = BOX_PACK_SIZES_BY_COMPONENT.get(sku)
        if pack_size:
            # Показуємо п'ять повних заводських упаковок, а поповнюємо,
            # коли лишається не більше двох упаковок.
            return pack_size * 5, pack_size * 2, band
        if is_cat_product(name):
            return CAT_CAN_TARGET, CAT_CAN_THRESHOLD, band
        if is_dog_product(name):
            return DOG_CAN_TARGET, DOG_CAN_THRESHOLD, band
    return DEFAULT_TARGET, DEFAULT_THRESHOLD, band


def base_skus(products, boxes):
    return set(products) - set(boxes)


def calculate_stock(sku, base_stock, boxes):
    if sku in boxes:
        component_sku, pack_size = boxes[sku]
        return max(0, int(base_stock.get(component_sku, 0)) // pack_size)
    return max(0, int(base_stock.get(sku, 0)))


def materialize_stock(products, base_stock, boxes):
    return {
        sku: calculate_stock(sku, base_stock, boxes)
        for sku in products
    }


def initial_base_stock(products, boxes, supplier):
    result = {}
    policies = {}
    for sku in base_skus(products, boxes):
        target, _, band = stock_policy(sku, products[sku], supplier)
        result[sku] = target
        policies[sku] = band
    return result, policies


def direct_order_snapshot(order, warehouse_id, products):
    snapshot = {}
    for item in order.get("products") or []:
        if int(item.get("stockId") or 0) != warehouse_id:
            continue

        sku = clean(item.get("sku"))
        if sku not in products:
            continue

        amount = int(float(item.get("amount") or 0))
        if amount > 0:
            snapshot[sku] = snapshot.get(sku, 0) + amount

    return snapshot


def find_warehouse_id(state, orders, products):
    configured = os.environ.get("SALESDRIVE_KORMAX_STOCK_ID", "").strip()
    if configured:
        return int(configured)
    if state and state.get("warehouse_id"):
        return int(state["warehouse_id"])

    for order in orders:
        if int(order.get("id") or 0) != BOOTSTRAP_ORDER_ID:
            continue
        for item in order.get("products") or []:
            sku = clean(item.get("sku"))
            stock_id = item.get("stockId")
            if sku in products and stock_id:
                return int(stock_id)

    raise RuntimeError(
        "Could not determine Kormax warehouse ID from order #2506. "
        "Set repository variable SALESDRIVE_KORMAX_STOCK_ID."
    )


def load_state():
    if not os.path.exists(STATE_FILE):
        return None
    try:
        with open(STATE_FILE, "r", encoding="utf-8") as file:
            state = json.load(file)
    except (OSError, json.JSONDecodeError):
        return None
    return state if state.get("version") == 1 else None


def save_state(state):
    os.makedirs(os.path.dirname(STATE_FILE), exist_ok=True)
    temporary = f"{STATE_FILE}.tmp"
    with open(temporary, "w", encoding="utf-8") as file:
        json.dump(state, file, ensure_ascii=False, indent=2, sort_keys=True)
        file.write("\n")
    os.replace(temporary, STATE_FILE)


def expand_consumption(sku, amount, boxes):
    if sku in boxes:
        component_sku, pack_size = boxes[sku]
        return {component_sku: amount * pack_size}
    return {sku: amount}


def apply_order_changes(products, boxes, warehouse_id, orders, state, supplier):
    base_stock = {}
    previous_policies = state.get("supplier_bands", {})

    for sku in base_skus(products, boxes):
        target, _, band = stock_policy(sku, products[sku], supplier)
        if sku in state.get("base_stock", {}):
            current = int(state["base_stock"][sku])
        else:
            current = target

        # Постачальник закінчив товар або знизив доступний діапазон —
        # одразу зменшуємо наш залишок до дозволеної межі.
        if target == 0:
            current = 0
        elif current > target:
            current = target
        # Товар повернувся у продаж після нульового залишку.
        elif previous_policies.get(sku) == "zero" and current == 0:
            current = target

        base_stock[sku] = current

    snapshots = state.get("order_snapshots", {})
    direct_deltas = {}

    for order in orders:
        order_id = str(order.get("id"))
        previous = snapshots.get(order_id, {})
        current = direct_order_snapshot(order, warehouse_id, products)

        for sku in set(previous) | set(current):
            delta = int(current.get(sku, 0)) - int(previous.get(sku, 0))
            if not delta:
                continue

            direct_deltas[sku] = direct_deltas.get(sku, 0) + delta
            for component_sku, component_delta in expand_consumption(
                sku, delta, boxes
            ).items():
                if component_sku in base_stock:
                    base_stock[component_sku] = max(
                        0,
                        base_stock[component_sku] - component_delta,
                    )

        snapshots[order_id] = current

    return base_stock, snapshots, direct_deltas


def replenish(products, boxes, supplier, base_stock):
    replenished = []
    policies = {}

    for sku in sorted(base_stock):
        target, threshold, band = stock_policy(sku, products[sku], supplier)
        policies[sku] = band

        if target == 0:
            base_stock[sku] = 0
        elif base_stock[sku] <= threshold:
            base_stock[sku] = target
            replenished.append(sku)

    return replenished, policies


def update_salesdrive_stock(api_key, warehouse_id, updates):
    items = [
        {
            "id": sku,
            "stockBalanceByStock": {str(warehouse_id): int(quantity)},
        }
        for sku, quantity in sorted(updates.items())
    ]

    for offset in range(0, len(items), 100):
        result = api_json(
            PRODUCT_UPDATE_URL,
            api_key,
            payload={"action": "update", "product": items[offset:offset + 100]},
        )
        if result.get("status") not in (None, "success"):
            raise RuntimeError(f"Could not update SalesDrive stock: {result}")


def build_yml(products, published_stock):
    root = etree.Element(
        "yml_catalog",
        date=now_kyiv().strftime("%Y-%m-%d %H:%M"),
    )
    shop = etree.SubElement(root, "shop")
    etree.SubElement(shop, "name").text = "CatPaws Kormax stock"
    etree.SubElement(shop, "company").text = "CatPaws"
    etree.SubElement(shop, "url").text = "https://catpaws.com.ua/"
    currencies = etree.SubElement(shop, "currencies")
    etree.SubElement(currencies, "currency", id="UAH", rate="1")
    categories = etree.SubElement(shop, "categories")
    etree.SubElement(categories, "category", id="1").text = "Kormax"
    offers = etree.SubElement(shop, "offers")

    for sku, name in sorted(products.items()):
        quantity = int(published_stock.get(sku, 0))
        offer = etree.SubElement(
            offers,
            "offer",
            id=sku,
            available="true" if quantity > 0 else "false",
        )
        etree.SubElement(offer, "name").text = name
        etree.SubElement(offer, "vendorCode").text = sku
        etree.SubElement(offer, "price").text = "1"
        etree.SubElement(offer, "currencyId").text = "UAH"
        etree.SubElement(offer, "categoryId").text = "1"
        etree.SubElement(offer, "quantity_in_stock").text = str(quantity)
        etree.SubElement(offer, "stock").text = str(quantity)
        etree.SubElement(offer, "in_stock").text = "1" if quantity > 0 else "0"

    return etree.ElementTree(root)


def main():
    api_key = os.environ.get("SALESDRIVE_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("SALESDRIVE_API_KEY is not configured")

    site_catalog = load_site_catalog()
    supplier = load_supplier()
    products, boxes = build_product_catalog(site_catalog, supplier)
    state = load_state()
    finished_at = now_kyiv()
    initializing = state is None

    if initializing:
        orders = fetch_orders(
            api_key,
            finished_at - timedelta(days=30),
            finished_at,
        )
        warehouse_id = find_warehouse_id(None, orders, products)
        base_stock, supplier_bands = initial_base_stock(
            products, boxes, supplier
        )
        snapshots = {
            str(order.get("id")): direct_order_snapshot(
                order, warehouse_id, products
            )
            for order in orders
        }
        published_stock = materialize_stock(products, base_stock, boxes)
        updates = published_stock
        replenished = []
        print("Initial Kormax stock synchronization")
    else:
        warehouse_id = find_warehouse_id(state, [], products)
        orders = fetch_orders(
            api_key,
            parse_api_time(state["last_sync"]) - timedelta(minutes=2),
            finished_at,
        )
        previous_published = {
            sku: int(quantity)
            for sku, quantity in state.get("published_stock", {}).items()
        }
        base_stock, snapshots, direct_deltas = apply_order_changes(
            products, boxes, warehouse_id, orders, state, supplier
        )
        replenished, supplier_bands = replenish(
            products, boxes, supplier, base_stock
        )
        published_stock = materialize_stock(products, base_stock, boxes)

        # SalesDrive уже списав безпосередньо замовлені SKU. Надсилаємо
        # лише коригування коробок, змін постачальника та поповнення порога.
        automatic_stock = dict(previous_published)
        for sku, delta in direct_deltas.items():
            automatic_stock[sku] = max(
                0,
                int(automatic_stock.get(sku, 0)) - delta,
            )

        updates = {
            sku: quantity
            for sku, quantity in published_stock.items()
            if int(automatic_stock.get(sku, 0)) != int(quantity)
        }

    if updates:
        update_salesdrive_stock(api_key, warehouse_id, updates)

    save_state({
        "version": 1,
        "warehouse_id": warehouse_id,
        "last_sync": format_api_time(finished_at),
        "base_stock": base_stock,
        "published_stock": published_stock,
        "supplier_bands": supplier_bands,
        "order_snapshots": snapshots,
    })

    os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
    build_yml(products, published_stock).write(
        OUTPUT_FILE,
        encoding="UTF-8",
        xml_declaration=True,
        pretty_print=True,
    )

    positive = sum(quantity > 0 for quantity in published_stock.values())
    print(f"Kormax supplier products selected: {len(supplier)}")
    print(f"Kormax products in CatPaws catalog: {len(products)}")
    print(f"Derived boxes: {len(boxes)}")
    print(f"SalesDrive warehouse ID: {warehouse_id}")
    print(f"Updated orders read: {len(orders)}")
    print(f"Products sent to SalesDrive: {len(updates)}")
    print(f"Products replenished by threshold: {len(replenished)}")
    print(f"Products with positive stock: {positive}")
    print(f"Products with zero stock: {len(products) - positive}")
    print(f"Created: {OUTPUT_FILE}")
    print(f"State saved: {STATE_FILE}")


if __name__ == "__main__":
    main()
